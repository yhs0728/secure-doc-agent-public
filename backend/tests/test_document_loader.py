from pathlib import Path

from docx import Document as DocxDocument

from ingestion import (
    document_loader as loader_module,
)
from ingestion.document_loader import (
    load_csv_document,
    load_docx_document,
    load_pdf_document,
    load_source_document,
)

import pytest

from ingestion.document_loader import (
    SourceDocumentLoadError,
    load_text_document,
)


def test_load_text_document(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "security-policy.txt"
    file_path.write_text(
        "사용자는 비밀번호를 90일마다 변경해야 한다.",
        encoding="utf-8",
    )

    document = load_text_document(
        file_path=file_path,
        document_id="doc-003",
        title="정보보안 운영 규정",
    )

    assert "90일" in document.page_content
    assert document.metadata["document_id"] == "doc-003"
    assert document.metadata["title"] == "정보보안 운영 규정"
    assert document.metadata["source"] == "security-policy.txt"


def test_load_text_document_raises_error_when_missing(
    tmp_path: Path,
) -> None:
    missing_file = tmp_path / "missing.txt"

    with pytest.raises(
        SourceDocumentLoadError,
        match="원본 문서를 찾을 수 없습니다",
    ):
        load_text_document(
            file_path=missing_file,
            document_id="doc-999",
            title="없는 문서",
        )


def test_load_text_document_rejects_empty_file(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "empty.txt"
    file_path.write_text(
        "   ",
        encoding="utf-8",
    )

    with pytest.raises(
        SourceDocumentLoadError,
        match="원본 문서가 비어 있습니다",
    ):
        load_text_document(
            file_path=file_path,
            document_id="doc-001",
            title="빈 문서",
        )
        
def test_load_csv_document(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "users.csv"

    file_path.write_text(
        (
            "name,department\n"
            "홍길동,보안팀\n"
            "김철수,인프라팀\n"
        ),
        encoding="utf-8",
    )

    document = load_csv_document(
        file_path=file_path,
        document_id="users",
        title="사용자 목록",
    )

    assert "name | department" in (
        document.page_content
    )
    assert "홍길동 | 보안팀" in (
        document.page_content
    )
    assert document.metadata["source"] == (
        "users.csv"
    )
    assert document.metadata["file_type"] == (
        ".csv"
    )
    
def test_load_docx_document(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "access-policy.docx"

    docx_document = DocxDocument()
    docx_document.add_heading(
        "접근 통제 정책",
        level=1,
    )
    docx_document.add_paragraph(
        "외부 접속에는 다중 인증이 필요합니다."
    )
    docx_document.save(
        str(file_path)
    )

    document = load_docx_document(
        file_path=file_path,
        document_id="access-policy",
        title="접근 통제 정책",
    )

    assert "접근 통제 정책" in (
        document.page_content
    )
    assert (
        "외부 접속에는 다중 인증이 필요합니다."
        in document.page_content
    )
    assert document.metadata["source"] == (
        "access-policy.docx"
    )
    assert document.metadata["file_type"] == (
        ".docx"
    )
    
def test_load_pdf_document(
    tmp_path: Path,
    monkeypatch,
) -> None:
    file_path = tmp_path / "security-policy.pdf"

    file_path.write_bytes(
        b"%PDF-1.4 test"
    )

    class FakePage:
        def extract_text(
            self,
        ) -> str:
            return (
                "비밀번호 변경 주기는 "
                "90일입니다."
            )

    class FakePdfReader:
        def __init__(
            self,
            file_name: str,
        ) -> None:
            assert file_name == str(file_path)
            self.pages = [
                FakePage(),
            ]

    monkeypatch.setattr(
        loader_module,
        "PdfReader",
        FakePdfReader,
    )

    document = load_pdf_document(
        file_path=file_path,
        document_id="security-policy",
        title="보안 정책",
    )

    assert document.page_content == (
        "비밀번호 변경 주기는 "
        "90일입니다."
    )
    assert document.metadata["source"] == (
        "security-policy.pdf"
    )
    assert document.metadata["file_type"] == (
        ".pdf"
    )
    
def test_load_source_document_routes_csv(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "policy.csv"

    file_path.write_text(
        (
            "policy,value\n"
            "password-change,90 days\n"
        ),
        encoding="utf-8",
    )

    document = load_source_document(
        file_path=file_path,
        document_id="policy",
        title="정책",
    )

    assert "password-change | 90 days" in (
        document.page_content
    )
    assert document.metadata["file_type"] == (
        ".csv"
    )