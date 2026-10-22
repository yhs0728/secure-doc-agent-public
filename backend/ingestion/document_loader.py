import csv
from pathlib import Path

from docx import Document as DocxDocument
from langchain_core.documents import (
    Document as LangChainDocument,
)
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".pdf",
    ".docx",
    ".csv",
}


class SourceDocumentLoadError(RuntimeError):
    """RAG 원본 문서를 읽지 못했을 때 발생합니다."""


def load_source_document(
    file_path: Path,
    document_id: str,
    title: str | None = None,
) -> LangChainDocument:
    extension = file_path.suffix.lower()

    if extension == ".txt":
        return load_text_document(
            file_path=file_path,
            document_id=document_id,
            title=title,
        )

    if extension == ".pdf":
        return load_pdf_document(
            file_path=file_path,
            document_id=document_id,
            title=title,
        )

    if extension == ".docx":
        return load_docx_document(
            file_path=file_path,
            document_id=document_id,
            title=title,
        )

    if extension == ".csv":
        return load_csv_document(
            file_path=file_path,
            document_id=document_id,
            title=title,
        )

    raise SourceDocumentLoadError(
        f"지원하지 않는 문서 형식입니다: {extension}"
    )


def build_langchain_document(
    content: str,
    file_path: Path,
    document_id: str,
    title: str | None = None,
) -> LangChainDocument:
    normalized_content = content.strip()

    if not normalized_content:
        raise SourceDocumentLoadError(
            f"원본 문서가 비어 있습니다: {file_path}"
        )

    return LangChainDocument(
        page_content=normalized_content,
        metadata={
            "document_id": document_id,
            "title": title or file_path.stem,
            "source": file_path.name,
            "file_type": file_path.suffix.lower(),
        },
    )


def load_text_document(
    file_path: Path,
    document_id: str,
    title: str | None = None,
) -> LangChainDocument:
    try:
        content = file_path.read_text(
            encoding="utf-8-sig",
        )

    except FileNotFoundError as error:
        raise SourceDocumentLoadError(
            f"원본 문서를 찾을 수 없습니다: {file_path}"
        ) from error

    except (
        UnicodeDecodeError,
        OSError,
    ) as error:
        raise SourceDocumentLoadError(
            f"TXT 문서를 읽을 수 없습니다: {file_path}"
        ) from error

    return build_langchain_document(
        content=content,
        file_path=file_path,
        document_id=document_id,
        title=title,
    )


def load_pdf_document(
    file_path: Path,
    document_id: str,
    title: str | None = None,
) -> LangChainDocument:
    try:
        reader = PdfReader(
            str(file_path)
        )

        page_contents = [
            page.extract_text() or ""
            for page in reader.pages
        ]

    except FileNotFoundError as error:
        raise SourceDocumentLoadError(
            f"원본 문서를 찾을 수 없습니다: {file_path}"
        ) from error

    except Exception as error:
        raise SourceDocumentLoadError(
            f"PDF 문서를 읽을 수 없습니다: {file_path}"
        ) from error

    return build_langchain_document(
        content="\n\n".join(page_contents),
        file_path=file_path,
        document_id=document_id,
        title=title,
    )


def load_docx_document(
    file_path: Path,
    document_id: str,
    title: str | None = None,
) -> LangChainDocument:
    try:
        document = DocxDocument(
            str(file_path)
        )

        contents = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip()
                    for cell in row.cells
                )

                if row_text.strip():
                    contents.append(row_text)

    except Exception as error:
        raise SourceDocumentLoadError(
            f"DOCX 문서를 읽을 수 없습니다: {file_path}"
        ) from error

    return build_langchain_document(
        content="\n".join(contents),
        file_path=file_path,
        document_id=document_id,
        title=title,
    )


def load_csv_document(
    file_path: Path,
    document_id: str,
    title: str | None = None,
) -> LangChainDocument:
    try:
        with file_path.open(
            "r",
            encoding="utf-8-sig",
            newline="",
        ) as csv_file:
            reader = csv.reader(csv_file)

            rows = [
                " | ".join(
                    cell.strip()
                    for cell in row
                )
                for row in reader
                if any(
                    cell.strip()
                    for cell in row
                )
            ]

    except FileNotFoundError as error:
        raise SourceDocumentLoadError(
            f"원본 문서를 찾을 수 없습니다: {file_path}"
        ) from error

    except (
        UnicodeDecodeError,
        csv.Error,
        OSError,
    ) as error:
        raise SourceDocumentLoadError(
            f"CSV 문서를 읽을 수 없습니다: {file_path}"
        ) from error

    return build_langchain_document(
        content="\n".join(rows),
        file_path=file_path,
        document_id=document_id,
        title=title,
    )


def load_source_document(
    file_path: Path,
    document_id: str,
    title: str | None = None,
) -> LangChainDocument:
    extension = file_path.suffix.lower()

    loaders = {
        ".txt": load_text_document,
        ".pdf": load_pdf_document,
        ".docx": load_docx_document,
        ".csv": load_csv_document,
    }

    loader = loaders.get(extension)

    if loader is None:
        raise SourceDocumentLoadError(
            f"지원하지 않는 문서 형식입니다: {extension}"
        )

    return loader(
        file_path=file_path,
        document_id=document_id,
        title=title,
    )