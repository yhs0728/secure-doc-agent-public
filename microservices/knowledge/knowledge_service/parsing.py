import csv
import io
import os

from docx import Document
from pypdf import PdfReader

SUPPORTED = {".txt", ".pdf", ".docx", ".csv"}


def parse_and_split(content, extension):
    if extension == ".pdf":
        pdf = PdfReader(io.BytesIO(content))
        pages = [(i + 1, page.extract_text() or "") for i, page in enumerate(pdf.pages)]
    elif extension == ".docx":
        doc = Document(io.BytesIO(content))
        lines = [p.text for p in doc.paragraphs]
        lines.extend(
            " | ".join(cell.text for cell in row.cells)
            for table in doc.tables
            for row in table.rows
        )
        pages = [(1, "\n".join(lines))]
    elif extension == ".csv":
        rows = csv.reader(io.StringIO(content.decode("utf-8-sig")))
        pages = [(1, "\n".join(" | ".join(row) for row in rows))]
    elif extension == ".txt":
        pages = [(1, content.decode("utf-8-sig"))]
    else:
        raise ValueError("UNSUPPORTED_FILE_TYPE")
    size = int(os.getenv("CHUNK_SIZE", "500"))
    overlap = int(os.getenv("CHUNK_OVERLAP", "50"))
    if size < 1 or not 0 <= overlap < size:
        raise ValueError("INVALID_CHUNK_CONFIGURATION")
    chunks = []
    for page, text in pages:
        text = text.strip()
        for start in range(0, len(text), size - overlap):
            chunk = text[start : start + size].strip()
            if chunk:
                chunks.append({"page": page, "text": chunk})
            if start + size >= len(text):
                break
    if not chunks:
        raise ValueError("EMPTY_DOCUMENT")
    return chunks, max(len(pages), 1)
